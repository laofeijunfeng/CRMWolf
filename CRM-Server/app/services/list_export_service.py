"""Streaming XLSX workbook writer for filtered list exports."""

from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from itertools import islice
from pathlib import Path
from re import compile as compile_regex
from tempfile import NamedTemporaryFile
from typing import TypeVar
from urllib.parse import quote

from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.cell import Cell as OpenpyxlCell
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from starlette.background import BackgroundTask

from app.core.list_export import ListExportCatalog, ListExportField
from app.utils.time import business_now

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_INVALID_FILE_CHARS = str.maketrans("", "", '\\/:*?"<>|\r\n')
_INVALID_SHEET_CHARS = compile_regex(r"[\\/*?:\[\]]")

T = TypeVar("T")


@dataclass(frozen=True)
class GeneratedListExport:
    path: Path
    filename: str
    row_count: int


def iter_batches(rows: Iterable[T], batch_size: int = 500) -> Iterator[list[T]]:
    """Slice a streaming iterator into fixed-size lists without draining it."""
    if batch_size < 1:
        raise ValueError("batch_size must be positive")
    iterator = iter(rows)
    while batch := list(islice(iterator, batch_size)):
        yield batch


def safe_excel_text(value: str) -> str:
    """Neutralize formula injection by checking the first non-whitespace character."""
    first = value.lstrip()[:1]
    return f"'{value}" if first in {"=", "+", "-", "@"} else value


def _safe_file_stem(value: str) -> str:
    cleaned = value.translate(_INVALID_FILE_CHARS).strip(" .")
    return cleaned or "列表导出"


def _safe_sheet_name(value: str) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("", value).strip("'")
    return (cleaned or "导出")[:31]

def _export_cell(sheet: WriteOnlyWorksheet, export_field: ListExportField, value: object) -> OpenpyxlCell:
    cell = WriteOnlyCell(sheet)
    if value is None:
        return cell
    if export_field.cell_type == "text":
        cell.value = safe_excel_text(str(value))
        cell.data_type = "s"
        return cell
    if export_field.cell_type in {"number", "currency"}:
        if not isinstance(value, (int, float, Decimal)) or isinstance(value, bool):
            raise TypeError(f"字段 {export_field.key} 必须是数值")
        cell.value = value
        cell.number_format = "#,##0.00" if export_field.cell_type == "currency" else "0.00"
        return cell
    if export_field.cell_type == "date":
        if not isinstance(value, date) or isinstance(value, datetime):
            raise TypeError(f"字段 {export_field.key} 必须是日期")
        cell.value = value
        cell.number_format = "yyyy-mm-dd"
        return cell
    if not isinstance(value, datetime):
        raise TypeError(f"字段 {export_field.key} 必须是日期时间")
    cell.value = value
    cell.number_format = "yyyy-mm-dd hh:mm:ss"
    return cell


def create_list_export_file(
    *,
    catalog: ListExportCatalog,
    selected_keys: Sequence[str],
    rows: Iterable[Mapping[str, object]],
    file_stem: str,
    directory: Path | None = None,
) -> GeneratedListExport:
    selected = catalog.select(selected_keys)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=_safe_sheet_name(catalog.sheet_name))
    sheet.freeze_panes = "A2"
    header: list[OpenpyxlCell] = []
    for export_field in selected:
        cell = WriteOnlyCell(sheet, value=safe_excel_text(export_field.label))
        cell.data_type = "s"
        cell.font = Font(bold=True)
        header.append(cell)
    sheet.append(header)

    row_count = 0
    for row in rows:
        sheet.append([
            _export_cell(sheet, export_field, export_field.value_from(row))
            for export_field in selected
        ])
        row_count += 1

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(selected))}{row_count + 1}"
    timestamp = business_now().strftime("%Y%m%d-%H%M%S")
    filename = f"{_safe_file_stem(file_stem)}-{timestamp}.xlsx"
    with NamedTemporaryFile(delete=False, suffix=".xlsx", dir=directory) as temporary:
        path = Path(temporary.name)
    try:
        workbook.save(path)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return GeneratedListExport(path=path, filename=filename, row_count=row_count)


def list_export_file_response(generated: GeneratedListExport) -> FileResponse:
    encoded = quote(generated.filename)
    return FileResponse(
        generated.path,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
        background=BackgroundTask(generated.path.unlink, missing_ok=True),
    )
