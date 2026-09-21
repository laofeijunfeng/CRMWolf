from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.core.list_export import ListExportCatalog, ListExportField, run_list_export_or_400
from app.services.list_export_service import (
    create_list_export_file,
    iter_batches,
    safe_excel_text,
)


def test_catalog_rejects_unsafe_and_unknown_fields() -> None:
    with pytest.raises(ValueError, match="不安全导出字段"):
        ListExportCatalog("bad", "Bad", (ListExportField("id", "ID", "text"),))

    catalog = ListExportCatalog(
        "customers",
        "客户",
        (ListExportField("account_name", "客户名称", "text"),),
    )
    with pytest.raises(ValueError, match="未知导出字段"):
        catalog.select(["missing"])


def test_http_helper_converts_catalog_error_to_400() -> None:
    catalog = ListExportCatalog(
        "customers",
        "客户",
        (ListExportField("account_name", "客户名称", "text"),),
    )
    with pytest.raises(HTTPException) as error:
        run_list_export_or_400(lambda: catalog.select(["missing"]))
    assert error.value.status_code == 400
    assert error.value.detail == "未知导出字段: missing"


def test_safe_excel_text_checks_prefix_after_leading_whitespace() -> None:
    assert safe_excel_text("=1+1") == "'=1+1"
    assert safe_excel_text("  @SUM(A1:A2)") == "'  @SUM(A1:A2)"
    assert safe_excel_text("Acme") == "Acme"


def test_iter_batches_never_loads_the_full_iterator() -> None:
    assert list(iter_batches(iter(range(5)), batch_size=2)) == [[0, 1], [2, 3], [4]]
    with pytest.raises(ValueError, match="batch_size must be positive"):
        list(iter_batches(iter(()), batch_size=0))


def test_writer_creates_typed_workbook_and_zero_row_workbook(tmp_path: Path) -> None:
    catalog = ListExportCatalog(
        "customers",
        "客户列表",
        (
            ListExportField("public_id", "业务 ID", "text"),
            ListExportField("amount", "金额", "currency"),
            ListExportField("signed_on", "签署日期", "date"),
            ListExportField("created_at", "创建时间", "datetime"),
        ),
    )
    generated = create_list_export_file(
        catalog=catalog,
        selected_keys=["public_id", "amount", "signed_on", "created_at"],
        rows=iter((
            {
                "public_id": "CUS-0001",
                "amount": 1200.5,
                "signed_on": date(2026, 9, 21),
                "created_at": datetime(2026, 9, 21, 14, 30),
            },
        )),
        file_stem="客户列表-所有客户",
        directory=tmp_path,
    )
    workbook = load_workbook(generated.path, read_only=True, data_only=False)
    values = list(workbook.active.values)
    assert values[0] == ("业务 ID", "金额", "签署日期", "创建时间")
    data_row = values[1]
    assert data_row[0] == "CUS-0001"
    assert data_row[1] == 1200.5
    # openpyxl reads date serials back as datetimes with a zero time component.
    assert data_row[2] == datetime(2026, 9, 21, 0, 0)
    assert data_row[3] == datetime(2026, 9, 21, 14, 30)
    assert generated.row_count == 1
    empty = create_list_export_file(
        catalog=catalog,
        selected_keys=["public_id"],
        rows=iter(()),
        file_stem="空列表",
        directory=tmp_path,
    )
    assert list(load_workbook(empty.path, read_only=True).active.values) == [("业务 ID",)]
